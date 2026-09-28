import zipfile, re, json, os
from pathlib import Path

src = Path(r"c:\Users\recyc\OneDrive\nexcommsys\EdgePoint Industrial Data Dictionary 230921.docx")
fixtures = Path(r"C:\Users\public\data\est-pc\st\fixtures")
fixtures.mkdir(parents=True, exist_ok=True)
dest_docx = fixtures / src.name
dest_txt = fixtures / "EdgePoint Industrial Data Dictionary 230921.extracted.txt"

data = src.read_bytes()
dest_docx.write_bytes(data)
print(f"Copied {len(data)} bytes to {dest_docx}")

from docx import Document
doc = Document(str(dest_docx))

lines = []
lines.append("=== PARAGRAPHS ===")
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t:
        lines.append(t)

lines.append("\n=== TABLES ===")
for ti, table in enumerate(doc.tables):
    lines.append(f"\n--- Table {ti+1} ({len(table.rows)} rows x {len(table.columns) if table.rows else 0} cols) ---")
    for ri, row in enumerate(table.rows):
        cells = [c.text.replace('\n',' ').strip() for c in row.cells]
        lines.append(" | ".join(cells))

text = "\n".join(lines)
dest_txt.write_text(text, encoding="utf-8")
print(f"Wrote {len(text)} chars to {dest_txt}")
print(f"Paragraphs with text: {sum(1 for p in doc.paragraphs if p.text.strip())}")
print(f"Tables: {len(doc.tables)}")
